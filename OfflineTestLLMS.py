
import docx
from flask import Flask, request, abort
from werkzeug.exceptions import BadRequest
from Class.Contract import Contract
from Class.LLMProcess import LLM
from Class.LLMProcess import labels
from LLMmodels import SparkAPI
from Class import Utils
if __name__ == "__main__":
    try:
        # new[ 中电信数智科技有限公司安徽分公司安徽省肥东危化品预警系统运维项目分包二 ]塔类租赁合同（6.14李芳审核稿）.docx
        # [ 中电信数智科技有限公司安徽分公司安徽省肥东危化品预警系统运维项目分包三 ]电路服务合同（6.docx
        # new巢湖市金盾实业集团有限公司智慧出行系统建设项目云网融合服务合同（5.30李芳审核稿）.docx
        # document = docx.Document("testfiles/test_total.docx")
        document = docx.Document("testfiles/[ 中电信数智科技有限公司安徽分公司安徽省肥东危化品预警系统运维项目分包三 ]电路服务合同（6.docx")
        # document = docx.Document("testfiles/new巢湖市金盾实业集团有限公司智慧出行系统建设项目云网融合服务合同（5.30李芳审核稿）.docx")
        contract = Contract()
        contract.get_text(document)
        '''
        llm = contract.match_with_llm() 
        #print(llm.split_text[0])
        prompt_input = llm.generate_prompt(label_list=labels, ori_contract=llm.split_text[0], pos_idx="", total_parts="")
        print(prompt_input)
        result = llm.llm_inference(prompt_input)
        #print(Utils.is_valid_json(result))
        '''
        contract.match_with_llm()
        result = contract.get_result()
        for error in result['errors']:
            print(error)
    except BadRequest:
        abort(400, description='Invalid file')
